
import csv
import io
import logging

from sqlalchemy.exc import IntegrityError
from pydantic import ValidationError as PydanticValidationError

from app.extensions import db
from app.schemas.catalog import PanelSchema, InverterSchema, BatterySchema, WireSchema
from app.services.catalog_service import CatalogService
from app.errors import ValidationError

logger = logging.getLogger(__name__)

RANGE_SCHEMAS = {
    'panels': PanelSchema,
    'inverters': InverterSchema,
    'batteries': BatterySchema,
    'wires': WireSchema,
}

NATURAL_KEYS = {
    'panels': ('nombre',),
    'inverters': ('nombre',),
    'batteries': ('nombre',),
    'wires': ('tipo', 'material', 'seccion', 'no_conductores'),
}

TEXT_EXTENSIONS = {'.csv', '.tsv'}
EXCEL_EXTENSIONS = {'.xlsx', '.xls'}


def _extension(filename):
    name = (filename or '').lower()
    dot = name.rfind('.')
    return name[dot:] if dot != -1 else ''


def _parse_text(content, extension):
    text = content.decode('utf-8-sig', errors='replace')
    if extension == '.tsv':
        delimiter = '\t'
    else:
        sample = text[:4096]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t').delimiter
        except csv.Error:
            delimiter = ','
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [row for row in reader]
    return rows


def _parse_excel(content):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError):
        raise ValidationError('El archivo Excel no se pudo leer.', code='import.unreadable')
    sheet = workbook.active
    rows = []
    for cells in sheet.iter_rows(values_only=True):
        rows.append(['' if c is None else c for c in cells])
    workbook.close()
    return rows


def _rows_from_file(filename, content):
    extension = _extension(filename)
    if extension in EXCEL_EXTENSIONS:
        grid = _parse_excel(content)
    else:
        grid = _parse_text(content, extension)
    grid = [row for row in grid if any(str(c).strip() for c in row)]
    if not grid:
        raise ValidationError('El archivo no contiene filas.', code='import.empty')
    header = [str(c).strip() for c in grid[0]]
    return header, grid[1:]


def _header_map(header, cfg):
    known = {field.lower(): field for field in cfg['fields']}
    mapping = {}
    for index, name in enumerate(header):
        field = known.get(name.strip().lower())
        if field is not None:
            mapping[index] = field
    return mapping


def _coerce(cfg, values):
    coerced = {}
    for field, raw in values.items():
        if raw in (None, ''):
            continue
        value = raw
        if isinstance(value, str):
            value = value.strip()
            if value == '':
                continue
        try:
            if field in cfg['numeric']:
                value = float(value)
            elif field in cfg['integer']:
                value = int(float(value))
        except (TypeError, ValueError):
            raise ValidationError(f'Campo numérico inválido: {field}')
        coerced[field] = value
    return coerced


def _validate_ranges(resource, values):
    schema = RANGE_SCHEMAS.get(resource)
    if schema is None:
        return
    payload = {k: v for k, v in values.items() if k in schema.model_fields}
    if not payload:
        return
    try:
        schema(**payload)
    except PydanticValidationError as exc:
        field = exc.errors()[0]['loc'][0] if exc.errors() else 'desconocido'
        raise ValidationError(f'Valor fuera de rango para el campo: {field}')


def _find_existing(model, catalog_id, resource, values):
    filters = {'catalog_id': catalog_id}
    for key in NATURAL_KEYS[resource]:
        if key not in values:
            return None
        filters[key] = values[key]
    return model.query.filter_by(**filters).first()


class EquipmentImportService:

    @staticmethod
    def run(resource, cfg, org_id, filename, content, catalog_id=None):
        header, data_rows = _rows_from_file(filename, content)
        mapping = _header_map(header, cfg)
        if not mapping:
            raise ValidationError(
                'Ninguna columna del archivo coincide con los campos esperados.',
                code='import.no_columns',
            )

        catalog = CatalogService.resolve_target_catalog(org_id, catalog_id)
        model = cfg['model']
        created = 0
        updated = 0
        errors = []

        for offset, cells in enumerate(data_rows):
            line = offset + 2
            values = {}
            for index, field in mapping.items():
                if index < len(cells):
                    values[field] = cells[index]
            try:
                coerced = _coerce(cfg, values)
                missing = [f for f in cfg['required'] if coerced.get(f) in (None, '')]
                if missing:
                    raise ValidationError('Faltan campos: ' + ', '.join(missing))
                _validate_ranges(resource, coerced)
                with db.session.begin_nested():
                    existing = _find_existing(model, catalog.id, resource, coerced)
                    if existing is not None:
                        for field, value in coerced.items():
                            setattr(existing, field, value)
                        updated += 1
                    else:
                        row_values = {**cfg['defaults'], **coerced}
                        row = model(**row_values, catalog_id=catalog.id)
                        if resource == 'inverters' and row.power_max is None:
                            row.power_max = row.power
                        db.session.add(row)
                        created += 1
            except ValidationError as exc:
                errors.append({'row': line, 'msg': exc.message})
            except IntegrityError:
                errors.append({'row': line, 'msg': 'Ya existe un equipo con ese nombre.'})

        db.session.commit()
        return {'created': created, 'updated': updated, 'errors': errors,
                'catalog_id': catalog.id, 'catalog_nombre': catalog.nombre}
